"""Tests for the Student.is_teacher flag: form, highlighting, filters,
teacher-portal ordering, and Excel import/export."""
import datetime
import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from django.utils.timezone import localdate, make_aware

from attendance.models import StudentAttendanceRecord
from core.models import CoursePayment, Student, StudentTeacherLink, Teacher, User


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(
            phone='02200000000', password='adminpass', role=User.Role.ADMIN)
        t_user = User.objects.create_user(
            phone='02211111111', password='tpass', role=User.Role.TEACHER)
        cls.teacher = Teacher.objects.create(
            user=t_user, full_name='معلم الكورس', is_course=True)

        # Alphabetically the teacher-student comes LAST, so ordering-first is meaningful.
        cls.normal_a = Student.objects.create(
            full_name='أحمد عادي', national_id='70000000000001', student_code='IT001')
        cls.normal_b = Student.objects.create(
            full_name='باسم عادي', national_id='70000000000002', student_code='IT002')
        cls.flagged = Student.objects.create(
            full_name='يوسف المعلم', national_id='70000000000003', student_code='IT003',
            is_teacher=True)
        for s in (cls.normal_a, cls.normal_b, cls.flagged):
            StudentTeacherLink.objects.create(teacher=cls.teacher, student=s)
            StudentAttendanceRecord.objects.create(
                student=s, date=datetime.date(2025, 8, 1),
                check_in_time=make_aware(datetime.datetime(2025, 8, 1, 8, 0)),
                original_teacher=cls.teacher, assigned_teacher=cls.teacher,
            )
            CoursePayment.objects.create(student=s, course=cls.teacher,
                                         year=2025, month=8)

    def setUp(self):
        self.client = Client()
        self.client.login(phone='02200000000', password='adminpass')

    def _rows(self, response):
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        return list(ws.iter_rows(values_only=True))


class StudentModelAndFormTests(_Base):

    def test_default_is_false(self):
        s = Student.objects.create(full_name='جديد', national_id='79999999999999')
        self.assertFalse(s.is_teacher)

    def test_edit_form_sets_flag(self):
        url = reverse('admin_portal:student_edit', args=[self.normal_a.pk])
        response = self.client.get(url)
        self.assertContains(response, 'name="is_teacher"')
        form = response.context['form']
        data = {k: v for k, v in form.initial.items() if v is not None and k != 'image'}
        data['is_teacher'] = 'on'
        self.client.post(url, data)
        self.normal_a.refresh_from_db()
        self.assertTrue(self.normal_a.is_teacher)


class AdminPagesTests(_Base):

    def test_student_list_highlight_and_filter(self):
        url = reverse('admin_portal:student_list')
        response = self.client.get(url)
        self.assertContains(response, 'class="row-is-teacher"', count=1)
        self.assertContains(response, 'badge-is-teacher')

        only = self.client.get(url, {'is_teacher': '1'})
        self.assertEqual([s.pk for s in only.context['page_obj']], [self.flagged.pk])
        excl = self.client.get(url, {'is_teacher': '0'})
        self.assertNotIn(self.flagged.pk, [s.pk for s in excl.context['page_obj']])
        self.assertEqual(excl.context['total_count'], 2)

    def test_invalid_filter_value_ignored(self):
        response = self.client.get(reverse('admin_portal:student_list'), {'is_teacher': 'x'})
        self.assertEqual(response.context['total_count'], 3)

    def test_student_detail_and_history_badge(self):
        r = self.client.get(reverse('admin_portal:student_detail', args=[self.flagged.pk]))
        self.assertContains(r, 'badge-is-teacher')
        r = self.client.get(reverse('admin_portal:student_history', args=[self.flagged.pk]))
        self.assertContains(r, 'badge-is-teacher')
        self.assertContains(r, 'class="row-is-teacher"')
        r = self.client.get(reverse('admin_portal:student_history', args=[self.normal_a.pk]))
        self.assertNotContains(r, 'class="row-is-teacher"')

    def test_attendance_records_highlight_and_filter(self):
        url = reverse('admin_portal:attendance_records')
        r = self.client.get(url)
        self.assertContains(r, 'class="row-is-teacher"', count=1)
        self.assertEqual(self.client.get(url, {'is_teacher': '1'}).context['student_total'], 1)
        self.assertEqual(self.client.get(url, {'is_teacher': '0'}).context['student_total'], 2)

    def test_course_roster_filter_and_highlight(self):
        url = reverse('admin_portal:course_roster', args=[self.teacher.pk])
        r = self.client.get(url)
        self.assertContains(r, 'class="row-is-teacher"', count=1)
        r = self.client.get(url, {'is_teacher': '1'})
        self.assertEqual([row['student'].pk for row in r.context['roster']], [self.flagged.pk])

    def test_payments_list_filter_and_highlight(self):
        url = reverse('admin_portal:payments_list')
        r = self.client.get(url)
        self.assertContains(r, 'class="row-is-teacher"', count=1)
        r = self.client.get(url, {'is_teacher': '0'})
        self.assertEqual(len(r.context['page_obj'].object_list), 2)

    def test_teacher_students_page_marks_rows(self):
        r = self.client.get(reverse('admin_portal:teacher_students', args=[self.teacher.pk]))
        self.assertContains(r, 'data-is-teacher="1"', count=1)
        self.assertContains(r, 'id="is-teacher-filter"')


class ExcelTests(_Base):

    def test_student_export_column_and_filter(self):
        url = reverse('admin_portal:student_export')
        rows = self._rows(self.client.get(url))
        idx = rows[0].index('is_teacher')
        values = {r[0]: r[idx] for r in rows[1:]}
        self.assertEqual(values['يوسف المعلم'], 'yes')
        self.assertEqual(values['أحمد عادي'], 'no')
        rows = self._rows(self.client.get(url, {'is_teacher': '1'}))
        self.assertEqual(len(rows), 2)

    def test_import_reads_flag(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['full_name', 'national_id', 'is_teacher'])
        ws.append(['مستورد معلم', '71000000000001', 'yes'])
        ws.append(['مستورد عادي', '71000000000002', ''])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile('s.xlsx', buf.getvalue())
        self.client.post(reverse('admin_portal:student_import'), {'excel_file': upload})
        self.assertTrue(Student.objects.get(national_id='71000000000001').is_teacher)
        self.assertFalse(Student.objects.get(national_id='71000000000002').is_teacher)

    def test_attendance_export_column_and_filter(self):
        url = reverse('admin_portal:attendance_export')
        rows = self._rows(self.client.get(url, {'tab': 'students'}))
        idx = rows[0].index('طالب معلم')
        self.assertEqual(sorted(r[idx] for r in rows[1:]), ['لا', 'لا', 'نعم'])
        rows = self._rows(self.client.get(url, {'tab': 'students', 'is_teacher': '1'}))
        self.assertEqual(len(rows), 2)

    def test_payments_export_column_and_filter(self):
        rows = self._rows(self.client.get(
            reverse('admin_portal:payments_export'), {'is_teacher': '1'}))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][rows[0].index('طالب معلم')], 'نعم')

    def test_teacher_students_export_column(self):
        rows = self._rows(self.client.get(
            reverse('admin_portal:teacher_students_export', args=[self.teacher.pk])))
        idx = rows[0].index('is_teacher')
        self.assertIn('yes', [r[idx] for r in rows[1:]])


class TeacherPortalTests(_Base):

    def setUp(self):
        self.client = Client()
        self.client.login(phone='02211111111', password='tpass')
        self.url = reverse('teacher_portal:dashboard')

    def test_teacher_students_listed_first(self):
        rows = self.client.get(self.url).context['student_rows']
        self.assertEqual(rows[0]['student'].pk, self.flagged.pk)

    def test_teacher_students_first_with_other_sorts(self):
        for sort in ('name_asc', 'name_desc', 'avg_rating_desc', 'absent_first'):
            rows = self.client.get(self.url, {'sort': sort}).context['student_rows']
            self.assertEqual(rows[0]['student'].pk, self.flagged.pk, sort)
        rows = self.client.get(self.url, {'sort': 'name_desc'}).context['student_rows']
        self.assertEqual([r['student'].pk for r in rows[1:]],
                         [self.normal_b.pk, self.normal_a.pk])

    def test_dashboard_filter_and_highlight(self):
        r = self.client.get(self.url)
        self.assertContains(r, 'class="row-is-teacher"', count=1)
        only = self.client.get(self.url, {'is_teacher': '1'}).context['student_rows']
        self.assertEqual([x['student'].pk for x in only], [self.flagged.pk])
        excl = self.client.get(self.url, {'is_teacher': '0'}).context['student_rows']
        self.assertEqual(len(excl), 2)

    def test_teacher_history_highlight(self):
        r = self.client.get(reverse('teacher_portal:student_history', args=[self.flagged.pk]))
        self.assertContains(r, ' row-is-teacher"')
        self.assertContains(r, 'badge-is-teacher')

    def test_export_column_order_and_filter(self):
        url = reverse('teacher_portal:export_attendance')
        rows = self._rows(self.client.get(url))
        self.assertEqual(rows[1][0], 'يوسف المعلم')
        self.assertEqual(rows[1][rows[0].index('طالب معلم')], 'نعم')
        rows = self._rows(self.client.get(url, {'is_teacher': '0'}))
        self.assertEqual(len(rows), 3)
